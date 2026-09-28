#!/usr/bin/env python3
"""Build a clean, readable Cisco IQ executive PowerPoint (no overlapping text)."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from copy import deepcopy

# Brand
BLUE = RGBColor(0x00, 0x97, 0xD0)
DARK = RGBColor(0x0D, 0x22, 0x3A)
NAVY = RGBColor(0x14, 0x2C, 0x47)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x2B, 0x3C)
MUTED = RGBColor(0x4A, 0x5C, 0x6E)
LIGHT = RGBColor(0xEE, 0xF3, 0xF8)
LINE = RGBColor(0xC9, 0xD5, 0xE0)
SOFT = RGBColor(0xDC, 0xEE, 0xF8)
GREEN = RGBColor(0x0F, 0x8A, 0x5F)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def rgb(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    rgb(s, color)
    return s


def textbox(slide, l, t, w, h):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    return box, tf


def clear_p(p):
    p.clear()


def style(run, size, bold=False, color=INK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_lines(tf, lines, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, space_after=8, space_before=0):
    """lines: list of str. First uses existing paragraph."""
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i == 0:
            clear_p(p)
        p.alignment = align
        p.space_before = Pt(space_before)
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        style(run, size, bold=bold, color=color)


def bullet_block(slide, l, t, w, h, items, size=17, color=INK):
    box, tf = textbox(slide, l, t, w, h)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i == 0:
            clear_p(p)
        p.level = 0
        p.space_before = Pt(4)
        p.space_after = Pt(10)
        # manual bullet for consistent rendering
        run = p.add_run()
        run.text = "•  " + item
        style(run, size, color=color)
    return box


def footer(slide, num, total=14, dark=False):
    c = RGBColor(0x8A, 0x9A, 0xAA) if not dark else RGBColor(0x7A, 0x8C, 0x9E)
    box, tf = textbox(slide, Inches(0.6), Inches(7.05), Inches(10), Inches(0.28))
    add_lines(tf, ["Cisco IQ Platform  |  Executive Briefing"], size=11, color=c, space_after=0)
    box2, tf2 = textbox(slide, Inches(11.5), Inches(7.05), Inches(1.3), Inches(0.28))
    add_lines(tf2, [f"{num} / {total}"], size=11, color=c, align=PP_ALIGN.RIGHT, space_after=0)


def content_chrome(slide, eyebrow, title, subtitle=None):
    """White slide with left accent + header. Content starts ~1.9\" if subtitle else ~1.55\"."""
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    rect(slide, 0, 0, Inches(0.14), SLIDE_H, BLUE)
    # top rule under header
    y_title = Inches(0.55)
    box, tf = textbox(slide, Inches(0.65), Inches(0.28), Inches(12), Inches(0.28))
    add_lines(tf, [eyebrow.upper()], size=12, bold=True, color=BLUE, space_after=0)

    box, tf = textbox(slide, Inches(0.65), y_title, Inches(12), Inches(0.55))
    add_lines(tf, [title], size=28, bold=True, color=DARK, space_after=0)

    content_top = Inches(1.35)
    if subtitle:
        box, tf = textbox(slide, Inches(0.65), Inches(1.15), Inches(12), Inches(0.4))
        add_lines(tf, [subtitle], size=15, color=MUTED, space_after=0)
        content_top = Inches(1.65)

    rect(slide, Inches(0.65), content_top - Inches(0.12), Inches(2.0), Inches(0.045), BLUE)
    return content_top + Inches(0.15)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    total = 14

    # ===== 1 TITLE =====
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, SLIDE_W, SLIDE_H, DARK)
    rect(s, 0, 0, Inches(0.18), SLIDE_H, BLUE)
    box, tf = textbox(s, Inches(0.9), Inches(1.9), Inches(11), Inches(0.35))
    add_lines(tf, ["CISCO CUSTOMER EXPERIENCE"], size=14, bold=True, color=BLUE, space_after=0)
    box, tf = textbox(s, Inches(0.9), Inches(2.4), Inches(11), Inches(0.9))
    add_lines(tf, ["Cisco IQ Platform"], size=44, bold=True, color=WHITE, space_after=0)
    box, tf = textbox(s, Inches(0.9), Inches(3.45), Inches(10.5), Inches(1.0))
    add_lines(
        tf,
        [
            "AI-powered front door to Cisco Support & Professional Services",
            "Executive briefing for platform overview and customer activation",
        ],
        size=18,
        color=RGBColor(0xC2, 0xD2, 0xDE),
        space_after=6,
    )
    rect(s, Inches(0.9), Inches(4.7), Inches(1.6), Inches(0.05), BLUE)
    box, tf = textbox(s, Inches(0.9), Inches(5.0), Inches(10), Inches(0.7))
    add_lines(
        tf,
        ["Plan  ·  Deploy  ·  Manage  ·  Secure  ·  Optimize", "Generally Available (SaaS)  ·  H2 FY2026"],
        size=15,
        color=RGBColor(0x9A, 0xB0, 0xC0),
        space_after=6,
    )
    footer(s, 1, total, dark=True)

    # ===== 2 AGENDA =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(s, "Agenda", "Session roadmap")
    agenda = [
        ("01", "The challenge", "Why reactive support models fail under AI-era complexity"),
        ("02", "Platform definition", "What Cisco IQ is — and what it consolidates"),
        ("03", "Four pillars", "Clarity, resilience, resolution, contextualized services"),
        ("04", "Capabilities", "Insights, assessments, AI support, services automation"),
        ("05", "Deployment & value", "SaaS / hybrid / air-gapped + measurable outcomes"),
        ("06", "Activation plan", "How entitled customers start realizing value"),
    ]
    for i, (num, title, desc) in enumerate(agenda):
        y = top + Inches(i * 0.78)
        box, tf = textbox(s, Inches(0.7), y, Inches(0.7), Inches(0.55))
        add_lines(tf, [num], size=20, bold=True, color=BLUE, space_after=0)
        box, tf = textbox(s, Inches(1.5), y, Inches(3.8), Inches(0.55))
        add_lines(tf, [title], size=18, bold=True, color=DARK, space_after=0)
        box, tf = textbox(s, Inches(5.5), y, Inches(7.0), Inches(0.55))
        add_lines(tf, [desc], size=16, color=MUTED, space_after=0)
        if i < len(agenda) - 1:
            rect(s, Inches(1.5), y + Inches(0.62), Inches(11), Inches(0.015), LINE)
    footer(s, 2, total)

    # ===== 3 CHALLENGE =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(
        s,
        "The Challenge",
        "Yesterday’s support model cannot keep up",
        "IT teams face fragmentation, skills pressure, and rising expectations.",
    )
    challenges = [
        ("Fragmented visibility", [
            "Multiple portals, tools, and APIs",
            "Blind spots across inventory",
            "Missed signals before impact",
        ]),
        ("Reactive firefighting", [
            "Work starts after disruption",
            "High cognitive load on engineers",
            "Limited time for innovation",
        ]),
        ("Context restart tax", [
            "Re-explaining environments",
            "TAC often starts from zero",
            "Slow case resolution cycles",
        ]),
        ("Aging risk exposure", [
            "~40% of top-targeted vulns hit EOL devices",
            "Misconfigurations drive outages",
            "Prioritization without evidence",
        ]),
    ]
    card_w = Inches(2.95)
    gap = Inches(0.18)
    for i, (title, bullets) in enumerate(challenges):
        x = Inches(0.65) + (card_w + gap) * i
        # card background only — text placed with padding, no overlapping shapes inside
        rect(s, x, top, card_w, Inches(4.55), LIGHT)
        rect(s, x, top, card_w, Inches(0.08), BLUE)
        box, tf = textbox(s, x + Inches(0.18), top + Inches(0.3), card_w - Inches(0.36), Inches(1.0))
        add_lines(tf, [title], size=16, bold=True, color=DARK, space_after=0)
        bullet_block(
            s,
            x + Inches(0.18),
            top + Inches(1.4),
            card_w - Inches(0.36),
            Inches(2.9),
            bullets,
            size=14,
            color=INK,
        )
    footer(s, 3, total)

    # ===== 4 WHAT IS =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(s, "Platform Definition", "What is Cisco IQ?")
    rect(s, Inches(0.65), top, Inches(12.0), Inches(1.35), SOFT)
    box, tf = textbox(s, Inches(0.9), top + Inches(0.25), Inches(11.5), Inches(0.95))
    add_lines(
        tf,
        [
            "Cisco IQ is an AI-powered digital platform that consolidates portals, tools, and APIs",
            "into one intelligent interface — the single front door to Cisco Customer Experience (CX).",
        ],
        size=17,
        color=DARK,
        space_after=4,
    )
    defs = [
        ("Unified CX experience", "Support and Professional Services in one AI interface"),
        ("Agentic intelligence", "Agents adapt to each customer’s unique environment"),
        ("40+ years of expertise", "Institutional Cisco knowledge encoded and continuously available"),
        ("Governed AI actions", "AI surfaces insights; humans decide; Cisco remains accountable"),
    ]
    for i, (t, d) in enumerate(defs):
        y = top + Inches(1.65) + Inches(i * 0.85)
        rect(s, Inches(0.65), y, Inches(0.1), Inches(0.7), BLUE)
        box, tf = textbox(s, Inches(1.0), y + Inches(0.05), Inches(11.5), Inches(0.3))
        add_lines(tf, [t], size=17, bold=True, color=DARK, space_after=0)
        box, tf = textbox(s, Inches(1.0), y + Inches(0.35), Inches(11.5), Inches(0.3))
        add_lines(tf, [d], size=15, color=MUTED, space_after=0)
    footer(s, 4, total)

    # ===== 5 QUOTE =====
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, SLIDE_W, SLIDE_H, DARK)
    rect(s, 0, 0, Inches(0.18), SLIDE_H, BLUE)
    box, tf = textbox(s, Inches(1.0), Inches(1.7), Inches(11), Inches(0.35))
    add_lines(tf, ["LEADERSHIP PERSPECTIVE"], size=13, bold=True, color=BLUE, space_after=0)
    box, tf = textbox(s, Inches(1.0), Inches(2.3), Inches(11.2), Inches(2.4))
    add_lines(
        tf,
        [
            "“Cisco IQ is our boldest step yet in reimagining how customers",
            "interact with Cisco — from planning and design to optimization",
            "and transformation.”",
        ],
        size=24,
        color=WHITE,
        space_after=8,
    )
    box, tf = textbox(s, Inches(1.0), Inches(5.1), Inches(11), Inches(0.35))
    add_lines(tf, ["Liz Centoni"], size=16, bold=True, color=BLUE, space_after=0)
    box, tf = textbox(s, Inches(1.0), Inches(5.5), Inches(11), Inches(0.35))
    add_lines(tf, ["EVP & Chief Customer Experience Officer, Cisco"], size=14, color=RGBColor(0xA0, 0xB2, 0xC2), space_after=0)
    footer(s, 5, total, dark=True)

    # ===== 6 PILLARS =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(s, "Strategic Framework", "Four pillars of Cisco IQ")
    pillars = [
        ("01  Landscape Clarity", "Real-time, benchmarked view of Cisco assets, configurations, and risks — including peer posture comparisons by industry and size."),
        ("02  Proactive Resilience", "Surfaces and ranks critical risks before business impact. Not hundreds of alerts — a prioritized action list with evidence and reasoning."),
        ("03  Rapid Resolution", "When a case opens, full environment context is already available. TAC briefs itself. Resolution timelines can move from hours toward minutes."),
        ("04  Contextualized Services", "Global expert experience is embedded into workflows across the lifecycle — helping projects land faster with less rework."),
    ]
    for i, (title, body) in enumerate(pillars):
        y = top + Inches(i * 1.15)
        rect(s, Inches(0.65), y, Inches(12.0), Inches(1.0), LIGHT)
        box, tf = textbox(s, Inches(0.9), y + Inches(0.12), Inches(11.5), Inches(0.32))
        add_lines(tf, [title], size=17, bold=True, color=BLUE, space_after=0)
        box, tf = textbox(s, Inches(0.9), y + Inches(0.48), Inches(11.5), Inches(0.45))
        add_lines(tf, [body], size=14, color=INK, space_after=0)
    footer(s, 6, total)

    # ===== 7 CAPABILITIES =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(s, "Platform Capabilities", "Core modules that deliver outcomes")
    caps = [
        ("Predictive Asset Insights", [
            "Unified inventory tracking",
            "Lifecycle milestones & LDoS planning",
            "Field notices & security advisories",
            "Real-time asset visibility",
        ]),
        ("Adaptive Assessments", [
            "Security & configuration checks",
            "Compliance / regulatory reviews",
            "Quantum readiness assessments",
            "Prioritized remediation plans",
        ]),
        ("AI-Powered Support", [
            "Consolidated cases & RMAs",
            "AI-assisted root cause analysis",
            "Environment-aware troubleshooting",
            "Faster time to resolution",
        ]),
        ("Services Automation", [
            "Digital requirements gathering",
            "Documentation automation",
            "Migration workflows",
            "Test, validation & governance",
        ]),
    ]
    card_w = Inches(2.95)
    gap = Inches(0.18)
    for i, (title, bullets) in enumerate(caps):
        x = Inches(0.65) + (card_w + gap) * i
        rect(s, x, top, card_w, Inches(4.55), LIGHT)
        # header bar — text ABOVE content area, not overlapping bullets
        rect(s, x, top, card_w, Inches(0.85), NAVY)
        box, tf = textbox(s, x + Inches(0.12), top + Inches(0.18), card_w - Inches(0.24), Inches(0.55))
        add_lines(tf, [title], size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, space_after=0)
        bullet_block(
            s,
            x + Inches(0.18),
            top + Inches(1.1),
            card_w - Inches(0.36),
            Inches(3.2),
            bullets,
            size=14,
            color=INK,
        )
    footer(s, 7, total)

    # ===== 8 LIFECYCLE =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(s, "Lifecycle Coverage", "One connected journey across the technology lifecycle")
    stages = [
        ("1. Plan", "Assess readiness, risks, requirements, and quantum/compliance posture before investment decisions."),
        ("2. Deploy", "Accelerate design and implementation with expert-aligned workflows and automation."),
        ("3. Manage", "Maintain live inventory, operational health, lifecycle milestones, and support entitlement context."),
        ("4. Secure", "Track advisories, configuration risk, and prioritized remediation with evidence."),
        ("5. Optimize", "Benchmark peer posture, automate improvements, and prove business outcomes."),
    ]
    for i, (title, body) in enumerate(stages):
        y = top + Inches(i * 0.9)
        box, tf = textbox(s, Inches(0.7), y, Inches(2.4), Inches(0.7))
        add_lines(tf, [title], size=18, bold=True, color=BLUE, space_after=0)
        box, tf = textbox(s, Inches(3.3), y, Inches(9.3), Inches(0.75))
        add_lines(tf, [body], size=15, color=INK, space_after=0)
        if i < len(stages) - 1:
            rect(s, Inches(3.3), y + Inches(0.72), Inches(9.3), Inches(0.015), LINE)
    footer(s, 8, total)

    # ===== 9 DEPLOYMENT =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(s, "Deployment Flexibility", "Fits enterprise security and operating models")
    modes = [
        ("SaaS", "Generally Available", [
            "Fastest time-to-value",
            "Continuously updated cloud experience",
            "Secure connectivity to Cisco services",
            "Best default for most customers",
        ]),
        ("On-Prem Tethered", "Controlled hybrid", [
            "Sensitive workloads stay local",
            "Connected for intelligence updates",
            "Balances control and freshness",
            "Fits regulated hybrid estates",
        ]),
        ("Air-Gapped", "Maximum isolation", [
            "Disconnected environment support",
            "Strong governance boundaries",
            "Designed for high-security zones",
            "No compromise on operating policy",
        ]),
    ]
    card_w = Inches(3.95)
    gap = Inches(0.2)
    for i, (title, tag, bullets) in enumerate(modes):
        x = Inches(0.65) + (card_w + gap) * i
        rect(s, x, top, card_w, Inches(4.55), LIGHT)
        rect(s, x, top, card_w, Inches(0.08), BLUE)
        box, tf = textbox(s, x + Inches(0.25), top + Inches(0.3), card_w - Inches(0.5), Inches(0.45))
        add_lines(tf, [title], size=20, bold=True, color=DARK, space_after=0)
        # tag as plain text, not overlapping pill on title
        box, tf = textbox(s, x + Inches(0.25), top + Inches(0.85), card_w - Inches(0.5), Inches(0.35))
        add_lines(tf, [tag], size=14, bold=True, color=BLUE, space_after=0)
        bullet_block(
            s,
            x + Inches(0.25),
            top + Inches(1.4),
            card_w - Inches(0.5),
            Inches(2.9),
            bullets,
            size=15,
            color=INK,
        )
    footer(s, 9, total)

    # ===== 10 OUTCOMES =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(s, "Business Outcomes", "What changes for customers")
    outcomes = [
        ("Hours → Minutes", "Case resolution accelerates because environment context is already loaded when support starts."),
        ("Reactive → Proactive", "Risks are ranked with evidence before outages, advisories, and compliance gaps become incidents."),
        ("Blind spots → Clarity", "One inventory and advisory view replaces fragmented portals and partial spreadsheets."),
        ("Opinions → Proof", "IT leaders get defensible data for infrastructure age, risk, and investment conversations."),
    ]
    for i, (title, body) in enumerate(outcomes):
        col = i % 2
        row = i // 2
        x = Inches(0.65) + Inches(col * 6.25)
        y = top + Inches(row * 2.25)
        rect(s, x, y, Inches(6.0), Inches(2.0), LIGHT)
        rect(s, x, y, Inches(0.12), Inches(2.0), GREEN if col == 0 else BLUE)
        box, tf = textbox(s, x + Inches(0.4), y + Inches(0.35), Inches(5.3), Inches(0.45))
        add_lines(tf, [title], size=20, bold=True, color=DARK, space_after=0)
        box, tf = textbox(s, x + Inches(0.4), y + Inches(0.95), Inches(5.3), Inches(0.8))
        add_lines(tf, [body], size=15, color=INK, space_after=0)
    footer(s, 10, total)

    # ===== 11 PROOF =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(s, "Proof Points", "Public signals and early customer impact")
    stats = [
        ("~90", "global customers already engaged with Cisco IQ"),
        ("40%", "of outages begin as misconfigurations that were not caught in time"),
        ("40%", "of top-targeted vulnerabilities impacted end-of-life devices"),
    ]
    for i, (num, label) in enumerate(stats):
        x = Inches(0.65) + Inches(i * 4.15)
        rect(s, x, top, Inches(3.95), Inches(2.5), NAVY)
        box, tf = textbox(s, x + Inches(0.25), top + Inches(0.4), Inches(3.45), Inches(0.8))
        add_lines(tf, [num], size=40, bold=True, color=BLUE, align=PP_ALIGN.CENTER, space_after=0)
        box, tf = textbox(s, x + Inches(0.25), top + Inches(1.35), Inches(3.45), Inches(0.9))
        add_lines(tf, [label], size=14, color=WHITE, align=PP_ALIGN.CENTER, space_after=4)
    rect(s, Inches(0.65), top + Inches(2.8), Inches(12.0), Inches(1.7), SOFT)
    box, tf = textbox(s, Inches(0.95), top + Inches(3.05), Inches(11.4), Inches(1.25))
    add_lines(
        tf,
        [
            "Customer example (Nestlé): Cisco IQ gave IT leaders defensible data to discuss aging",
            "infrastructure with business stakeholders — turning negotiation without evidence into a",
            "conversation with proof.",
        ],
        size=15,
        color=DARK,
        space_after=4,
    )
    footer(s, 11, total)

    # ===== 12 POSITIONING =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(s, "Positioning", "Cisco IQ vs. prior CX experience")
    # table header
    rect(s, Inches(0.65), top, Inches(12.0), Inches(0.6), NAVY)
    box, tf = textbox(s, Inches(0.85), top + Inches(0.15), Inches(2.4), Inches(0.35))
    add_lines(tf, ["Dimension"], size=14, bold=True, color=WHITE, space_after=0)
    box, tf = textbox(s, Inches(3.4), top + Inches(0.15), Inches(4.4), Inches(0.35))
    add_lines(tf, ["Previous CX model"], size=14, bold=True, color=WHITE, space_after=0)
    box, tf = textbox(s, Inches(8.0), top + Inches(0.15), Inches(4.4), Inches(0.35))
    add_lines(tf, ["Cisco IQ"], size=14, bold=True, color=BLUE, space_after=0)

    rows = [
        ("Scope", "Support-focused portal", "Support + Professional Services, AI-native"),
        ("Operating model", "Often reactive / ticket-centric", "Proactive, predictive, personalized"),
        ("Intelligence", "Tool-assisted workflows", "Agentic AI adapted to environment"),
        ("Access", "Multiple portals & tools", "One front door to CX resources"),
        ("Business value", "Resolve after impact", "Prevent, prioritize, and prove outcomes"),
    ]
    for i, (a, b, c) in enumerate(rows):
        y = top + Inches(0.6) + Inches(i * 0.78)
        bg = LIGHT if i % 2 == 0 else WHITE
        rect(s, Inches(0.65), y, Inches(12.0), Inches(0.78), bg)
        box, tf = textbox(s, Inches(0.85), y + Inches(0.22), Inches(2.4), Inches(0.4))
        add_lines(tf, [a], size=14, bold=True, color=DARK, space_after=0)
        box, tf = textbox(s, Inches(3.4), y + Inches(0.22), Inches(4.4), Inches(0.4))
        add_lines(tf, [b], size=14, color=MUTED, space_after=0)
        box, tf = textbox(s, Inches(8.0), y + Inches(0.22), Inches(4.4), Inches(0.4))
        add_lines(tf, [c], size=14, bold=True, color=INK, space_after=0)
    footer(s, 12, total)

    # ===== 13 CTA =====
    s = prs.slides.add_slide(blank)
    top = content_chrome(
        s,
        "Call to Action",
        "Start with the platform you already have",
        "For existing Cisco Support & Professional Services customers, Cisco IQ is included — not a separate budget ask.",
    )
    steps = [
        ("1", "Activate access", "Enable Cisco IQ SaaS for entitled accounts and assign clear admin ownership."),
        ("2", "Connect inventory", "Bring assets and telemetry into landscape clarity for immediate visibility."),
        ("3", "Run assessments", "Prioritize security, configuration, compliance, and EOL exposures with evidence."),
        ("4", "Operationalize support", "Use AI-assisted cases so TAC starts with full environment context."),
    ]
    for i, (n, title, body) in enumerate(steps):
        y = top + Inches(i * 1.05)
        # number circle area as simple colored square to avoid text-in-oval clipping
        rect(s, Inches(0.65), y, Inches(0.65), Inches(0.85), BLUE)
        box, tf = textbox(s, Inches(0.65), y + Inches(0.22), Inches(0.65), Inches(0.45))
        add_lines(tf, [n], size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, space_after=0)
        rect(s, Inches(1.45), y, Inches(11.2), Inches(0.85), LIGHT)
        box, tf = textbox(s, Inches(1.7), y + Inches(0.1), Inches(10.7), Inches(0.3))
        add_lines(tf, [title], size=17, bold=True, color=DARK, space_after=0)
        box, tf = textbox(s, Inches(1.7), y + Inches(0.42), Inches(10.7), Inches(0.35))
        add_lines(tf, [body], size=14, color=MUTED, space_after=0)
    footer(s, 13, total)

    # ===== 14 CLOSE =====
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, SLIDE_W, SLIDE_H, DARK)
    rect(s, 0, 0, Inches(0.18), SLIDE_H, BLUE)
    box, tf = textbox(s, Inches(1.0), Inches(2.0), Inches(11), Inches(0.35))
    add_lines(tf, ["SUMMARY"], size=13, bold=True, color=BLUE, space_after=0)
    box, tf = textbox(s, Inches(1.0), Inches(2.5), Inches(11.2), Inches(0.8))
    add_lines(tf, ["From firefighting to foresight."], size=36, bold=True, color=WHITE, space_after=0)
    box, tf = textbox(s, Inches(1.0), Inches(3.5), Inches(11), Inches(1.1))
    add_lines(
        tf,
        [
            "Cisco IQ turns 40 years of Cisco expertise into a continuous,",
            "proactive, AI-powered customer experience.",
        ],
        size=18,
        color=RGBColor(0xC2, 0xD2, 0xDE),
        space_after=6,
    )
    rect(s, Inches(1.0), Inches(4.9), Inches(1.6), Inches(0.05), BLUE)
    box, tf = textbox(s, Inches(1.0), Inches(5.2), Inches(11), Inches(0.7))
    add_lines(
        tf,
        ["Cisco IQ  ·  One front door to Customer Experience", "Next step: iq.cisco.com"],
        size=16,
        color=WHITE,
        space_after=6,
    )
    footer(s, 14, total, dark=True)

    out = "/workspace/presentations/Cisco_IQ_Platform_Executive_Briefing.pptx"
    prs.save(out)
    return out, len(prs.slides)


if __name__ == "__main__":
    path, n = build()
    print(f"Wrote {path} ({n} slides)")
